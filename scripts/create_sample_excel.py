"""Generate sample_balance_sheet.xlsx from sample_balance_sheet.csv."""

import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(__file__), "..", "data")
CSV_PATH = os.path.join(BASE, "sample_balance_sheet.csv")
XLSX_PATH = os.path.join(BASE, "sample_balance_sheet.xlsx")


def main() -> None:
    df = pd.read_csv(CSV_PATH)

    instructions = pd.DataFrame([
        ["name", "Yes", "Instrument label", "Fixed Gov Bonds (5Y)"],
        ["side", "Yes", "asset or liability", "asset"],
        ["notional", "Yes", "Face value in USD millions", "300"],
        ["coupon_pct", "Yes", "Annual interest rate (%)", "3.95"],
        [
            "instrument_type",
            "Yes",
            "bullet_fixed | bullet_floating | amortising | demand_deposit",
            "bullet_fixed",
        ],
        ["maturity_years", "Yes", "Maturity in years (0.25 = 3 months)", "5.0"],
        [
            "payment_freq",
            "No",
            "Payments per year: 1=annual, 2=semi, 4=quarterly, 12=monthly",
            "2",
        ],
        [
            "repricing_years",
            "No",
            "Years to next rate reset (floating/NMD only)",
            "0.5",
        ],
    ], columns=["Column", "Required", "Description", "Example"])

    tenor_ref = pd.DataFrame([
        ["Overnight", 0.00274, "1/365"],
        ["1 month", 0.083, "1/12"],
        ["3 months", 0.25, ""],
        ["6 months", 0.5, ""],
        ["1 year", 1.0, ""],
        ["5 years", 5.0, ""],
        ["10 years", 10.0, ""],
    ], columns=["Tenor", "maturity_years value", "Formula"])

    with pd.ExcelWriter(XLSX_PATH, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Balance Sheet", index=False)
        instructions.to_excel(writer, sheet_name="Instructions", index=False)
        tenor_ref.to_excel(writer, sheet_name="Tenor Reference", index=False)

    wb = load_workbook(XLSX_PATH)
    header_fill = PatternFill("solid", fgColor="1A3A6B")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="D0D5E8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            max_len = max(len(str(cell.value or "")) for cell in ws[letter])
            ws.column_dimensions[letter].width = min(max_len + 3, 45)
        ws.freeze_panes = "A2"

    wb.save(XLSX_PATH)


if __name__ == "__main__":
    main()
