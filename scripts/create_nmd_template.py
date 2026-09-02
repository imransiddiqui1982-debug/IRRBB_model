"""Create customer-level NMD Excel template (3 years of EOM balances)."""

import os

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(__file__), "..", "data")
XLSX_PATH = os.path.join(BASE, "nmd_deposit_template.xlsx")

MONTHS = pd.date_range("2022-01-31", periods=36, freq="ME")
MONTH_COLS = [d.strftime("%Y-%m") for d in MONTHS]


def _customer_path(start_idx: int, start_bal: float, decay: float, noise: float, rng) -> list[float]:
    vals = []
    bal = start_bal
    for i in range(36):
        if i < start_idx:
            vals.append(np.nan)
        else:
            if i > start_idx:
                bal = max(bal * (1 - decay) + rng.normal(0, noise), 0.0)
            vals.append(round(bal, 4))
    return vals


def build_sample_customers(rng: np.random.Generator) -> pd.DataFrame:
    rows = []
    # Retail transactional — sticky
    for i in range(1, 9):
        start_idx = rng.integers(0, 6)
        path = _customer_path(start_idx, rng.uniform(0.8, 2.5), 0.008, 0.02, rng)
        rows.append({
            "customer_id": f"RT-{i:03d}",
            "segment": "retail_transactional",
            "start_date": MONTHS[start_idx].strftime("%Y-%m-%d"),
            "deposit_rate": round(float(rng.uniform(0.004, 0.012)), 4),
            **{m: v for m, v in zip(MONTH_COLS, path)},
        })
    # Retail non-transactional
    for i in range(1, 6):
        start_idx = rng.integers(0, 10)
        path = _customer_path(start_idx, rng.uniform(1.0, 4.0), 0.015, 0.05, rng)
        rows.append({
            "customer_id": f"RN-{i:03d}",
            "segment": "retail_non_transactional",
            "start_date": MONTHS[start_idx].strftime("%Y-%m-%d"),
            "deposit_rate": round(float(rng.uniform(0.010, 0.025)), 4),
            **{m: v for m, v in zip(MONTH_COLS, path)},
        })
    # Wholesale — more volatile
    for i in range(1, 4):
        start_idx = rng.integers(0, 8)
        path = _customer_path(start_idx, rng.uniform(5.0, 15.0), 0.025, 0.2, rng)
        rows.append({
            "customer_id": f"WS-{i:03d}",
            "segment": "wholesale",
            "start_date": MONTHS[start_idx].strftime("%Y-%m-%d"),
            "deposit_rate": round(float(rng.uniform(0.020, 0.040)), 4),
            **{m: v for m, v in zip(MONTH_COLS, path)},
        })
    return pd.DataFrame(rows)


def build_market_rates(rng: np.random.Generator) -> pd.DataFrame:
    rate = 0.045
    rates = []
    for d in MONTHS:
        rate = float(np.clip(rate + rng.normal(0, 0.0015), 0.01, 0.10))
        rates.append({"date": d.strftime("%Y-%m-%d"), "market_rate": round(rate, 5)})
    return pd.DataFrame(rates)


def style_workbook(path: str) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1A3A6B")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Side(style="thin", color="D0D5E8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.border = border
        for col in range(1, ws.max_column + 1):
            letter = get_column_letter(col)
            max_len = max(len(str(cell.value or "")) for cell in ws[letter])
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 18)
        ws.freeze_panes = "E2" if sheet_name == "Customer Balances" else "A2"
    wb.save(path)


def main() -> None:
    os.makedirs(BASE, exist_ok=True)
    rng = np.random.default_rng(42)
    customers = build_sample_customers(rng)
    market = build_market_rates(rng)

    instructions = pd.DataFrame([
        ["customer_id", "Yes", "Unique customer / account ID", "RT-001"],
        ["segment", "Yes", "retail_transactional | retail_non_transactional | wholesale", "retail_transactional"],
        ["start_date", "Yes", "Date the deposit relationship started", "2022-03-15"],
        ["deposit_rate", "Yes", "Rate paid to customer (decimal: 0.01 = 1%)", "0.008"],
        ["YYYY-MM columns", "Yes", "36 end-of-month balances in USD millions (leave blank before start)", "1.25"],
        ["Market Rates!date", "Yes", "Month-end date for market rate", "2022-01-31"],
        ["Market Rates!market_rate", "Yes", "Market / policy rate (decimal)", "0.045"],
    ], columns=["Field", "Required", "Description", "Example"])

    segment_caps = pd.DataFrame([
        ["retail_transactional", "90%", "5.0 years"],
        ["retail_non_transactional", "70%", "4.5 years"],
        ["wholesale", "50%", "4.0 years"],
    ], columns=["Segment", "Max Core % (BCBS 368)", "Max WAL"])

    with pd.ExcelWriter(XLSX_PATH, engine="openpyxl") as writer:
        customers.to_excel(writer, sheet_name="Customer Balances", index=False)
        market.to_excel(writer, sheet_name="Market Rates", index=False)
        instructions.to_excel(writer, sheet_name="Instructions", index=False)
        segment_caps.to_excel(writer, sheet_name="Segment Caps", index=False)

    style_workbook(XLSX_PATH)


if __name__ == "__main__":
    main()
